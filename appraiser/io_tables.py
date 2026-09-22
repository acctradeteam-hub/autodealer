"""Чтение табличных входов (CSV, при наличии openpyxl — XLSX) и разбор значений."""
from __future__ import annotations

import csv
import re
from pathlib import Path

MONEY_RE = re.compile(r"[^0-9,.\-]")


def _normalize_key(key: str) -> str:
    return re.sub(r"[\s\-]+", "_", (key or "").strip().lower())


def read_table(path: str | Path, aliases: dict[str, str] | None = None) -> list[dict]:
    """Читает CSV или XLSX в список словарей с нормализованными именами колонок.

    aliases — карта «как в вашем файле» -> «как в оценщике», например
    {"Odometer": "mileage"}. Ключи сравниваются без учёта регистра и пробелов.
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Файл не найден: {p}")
    if p.suffix.lower() in {".xlsx", ".xlsm"}:
        rows = _read_xlsx(p)
    else:
        with p.open(newline="", encoding="utf-8-sig") as fh:
            rows = [dict(row) for row in csv.DictReader(fh)]

    alias_map = {_normalize_key(k): v for k, v in (aliases or {}).items()}
    out: list[dict] = []
    for row in rows:
        clean: dict[str, str] = {}
        for key, value in row.items():
            if key is None:
                continue
            norm = _normalize_key(key)
            clean[alias_map.get(norm, norm)] = (value or "").strip() if isinstance(value, str) else value
        if any(str(v).strip() for v in clean.values()):
            out.append(clean)
    return out


def _read_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - зависит от окружения
        raise RuntimeError(
            f"Для чтения {path.name} нужен openpyxl: pip install openpyxl "
            "(или сохраните файл как CSV)"
        ) from exc
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(c) if c is not None else "" for c in next(rows)]
    except StopIteration:
        return []
    result = []
    for raw in rows:
        result.append({header[i]: raw[i] for i in range(min(len(header), len(raw)))})
    wb.close()
    return result


def to_float(value, default=None):
    """'$12,500.00' -> 12500.0; пустое значение -> default."""
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = MONEY_RE.sub("", str(value)).replace(",", "")
    if text in {"", "-", "."}:
        return default
    try:
        return float(text)
    except ValueError:
        return default


def to_int(value, default=None):
    result = to_float(value, None)
    return default if result is None else int(round(result))


def norm_text(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def model_key(make, model) -> str:
    return f"{norm_text(make)}|{norm_text(model)}"
